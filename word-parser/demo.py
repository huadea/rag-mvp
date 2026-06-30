"""
最简示例：创建一个测试 .docx，然后解析打印结果。
运行：uv run python demo.py
"""

from docx import Document
from pathlib import Path


# ── 1. 创建示例 .docx ──────────────────────────────────────────────────────────

def make_sample_docx(path: str):
    doc = Document()
    doc.add_heading("招标公告", level=1)
    doc.add_heading("项目基本信息", level=2)
    doc.add_paragraph("项目编号：ZB-2024-001")
    doc.add_paragraph("项目名称：某市政府办公楼采购项目")
    doc.add_paragraph("招标方式：公开招标")

    doc.add_heading("时间安排", level=2)
    doc.add_paragraph("投标截止时间：2024年3月15日 17:00")
    doc.add_paragraph("开标时间：2024年3月16日 09:00")
    doc.add_paragraph("开标地点：某市公共资源交易中心三楼开标室")

    doc.add_heading("标包信息", level=2)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    tbl.rows[0].cells[0].text = "标包号"
    tbl.rows[0].cells[1].text = "标包名称"
    tbl.rows[0].cells[2].text = "最高限价"
    row = tbl.add_row()
    row.cells[0].text = "包1"
    row.cells[1].text = "办公设备采购"
    row.cells[2].text = "50万元"
    row = tbl.add_row()
    row.cells[0].text = "包2"
    row.cells[1].text = "网络设备采购"
    row.cells[2].text = "30万元"

    doc.save(path)
    print(f"已生成示例文件：{path}\n")


# ── 2. 最简解析 ────────────────────────────────────────────────────────────────

def simple_parse(path: str):
    doc = Document(path)

    print("=== 段落 ===")
    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        prefix = "#" * int(para.style.name[-1]) + " " if "Heading" in para.style.name else ""
        print(f"{prefix}{para.text}")

    print("\n=== 表格 ===")
    for i, table in enumerate(doc.tables):
        print(f"-- 表格 {i+1} --")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            print(" | ".join(cells))


# ── 3. 运行 ────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    sample = "sample.docx"
    make_sample_docx(sample)
    simple_parse(sample)
    Path(sample).unlink()  # 演示完删除临时文件
